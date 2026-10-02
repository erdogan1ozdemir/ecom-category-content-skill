#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Marka başına brief Excel'ini kurar ve sayfa satırlarını doldurur.

Excel, içerik yazılacak tüm hedef sayfaları sayfa tipine göre sekmelere ayrılmış olarak taşır. Sekmeler
profilden (`sekmeler`) gelir; yoksa envanterden kurulur: "Kategori" + sitede cinsiyetli kategori varsa
cinsiyet sekmeleri (Kadın, Erkek, Çocuk, Bebek) + "Marka". Her sayfa bir satırdır; "Brief" sütunu başlangıçta
"Bekliyor"dur. Bir sayfanın briefi hazırlandığında ilgili satır doldurulur ve "Hazır" olur.

Kullanım:
    python3 brief_satiri.py --marka X --kur                       # sekmeleri envanterden kur (tüm kategori sayfaları)
    python3 brief_satiri.py --marka X --kur --hedefler liste.txt  # yalnız verilen adresler (satır başına bir URL)
    python3 brief_satiri.py --marka X --json satir.json           # satırı doldur
    python3 brief_satiri.py --marka X --ozet                      # sekme başına hazır / bekleyen

Dosya: --xlsx verilmezse profildeki `cikti_klasoru` içinde "{Marka} kategori içerik briefleri.xlsx"
(profilde `brief_dosyasi` varsa o ad). --kur var olan dosyada doldurulmuş satırları korur; yalnız eksik sayfaları
ekler. Envanterde bulunmayan bir URL'nin briefi gelirse satır ilgili sekmenin sonuna eklenir.
Kur kapsamı profilden: `kur_kaynaklar` (yalnız bu sitemap adları), `kur_haric` (adres/slug regex'i),
`kur_marka` (marka sayfaları da eklensin mi; varsayılan hayır, brief hazırlandıkça satır açılır).

satir.json biçimi (çok satırlı metinler "\\n" ile):
{
  "kategori": "Kadın Mont",
  "url": "https://www.boyner.com.tr/kadin-mont-x-g3731-c23896554",
  "main_kw": "kadın mont",
  "hacim": 33100,
  "ikincil": "kelime (hacim)\\nkelime (hacim)",
  "basliklar": "H2: ...\\nH3: ...",
  "kurgu": "TON: ...\\n\\nAÇILIŞ: ...",
  "kurgu_kalin": "(isteğe bağlı) DİKKAT'in sonuna kalın basılacak koşullu uyarı",
  "link": "1. anchor : https://...\\n   yerleşeceği bölüm ve cümle",
  "kapsam_disi": "kelime (hacim) -> sahibi olan sayfa",
  "sss": "1. Soru?\\n   Yanıtta: ...",
  "yanit": "• Her yanıt ...",
  "durum": "Sayfa tipi, ürün sayısı, mevcut içerik, sıralama, GSC özeti"
}
"""
import argparse, json, math, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ortak
from ortak import url_coz

INK, HEAD, FN = "FF10332F", "434343", "Calibri"
SEKME_SIRA = ["Kategori", "Kadın", "Erkek", "Çocuk", "Bebek", "Unisex", "Marka"]
BASLIK = ["Kategori", "URL", "Brief", "Main KW", "Main KW Hacim", "İkincil ve Uzun Kuyruk Kelimeler", "Alt Başlıklar",
          "İçerik Kurgusu", "Link Verilecek Sayfalar", "Kapsam Dışı Kelimeler (Sahibi Başka Sayfa)", "SSS'ler",
          "Yanıt Biçimi", "Mevcut Durum"]
GENISLIK = [26, 44, 11, 18, 12, 38, 48, 130, 72, 62, 66, 56, 44]
ANAHTAR = ["kategori", "url", "brief", "main_kw", "hacim", "ikincil", "basliklar", "kurgu", "link", "kapsam_disi",
           "sss", "yanit", "durum"]
ZORUNLU = [k for k in ANAHTAR if k != "brief"]
ORTALI = {"hacim", "brief"}


def sekmeler():
    if ortak.AYAR.get("sekmeler"):
        return list(ortak.AYAR["sekmeler"])
    var = {"Kategori", "Marka"}
    try:
        import envanter
        for s in envanter.yukle():
            if s["tip"].startswith("cinsiyet_kategori") and s.get("g"):
                var.add(ortak.sekme_adi_cinsiyet(s["g"]) or "Kategori")
    except SystemExit:
        pass
    return [x for x in SEKME_SIRA if x in var] + sorted(var - set(SEKME_SIRA))


def sekme_adi(url, mevcut):
    p = url_coz(url) or {}
    if p.get("b"):
        return "Marka" if "Marka" in mevcut else mevcut[-1]
    ad = ortak.sekme_adi_cinsiyet(p.get("g"))
    return ad if ad in mevcut else "Kategori"


def satir_sayisi(metin, genislik):
    kap = max(1, int(genislik * 1.15))
    return sum(max(1, math.ceil(len(p) / kap)) for p in str(metin).split("\n"))


def sekme_kur(wb, ad):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(ad)
    for c, b in enumerate(BASLIK, start=1):
        h = ws.cell(row=1, column=c, value=b)
        h.font = Font(name=FN, size=11, bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=HEAD)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = GENISLIK[c - 1]
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(BASLIK))}1"
    return ws


def hucre(ws, r, c, deger, k):
    from openpyxl.styles import Font, Alignment
    cell = ws.cell(row=r, column=c, value=deger)
    cell.font = Font(name=FN, size=10, color=INK)
    cell.alignment = Alignment(horizontal="center" if k in ORTALI else "left", vertical="center", wrap_text=True)
    return cell


def ac(yol):
    from openpyxl import Workbook, load_workbook
    if os.path.exists(yol):
        wb = load_workbook(yol, rich_text=True)
    else:
        wb = Workbook(); wb.remove(wb.active)
    for ad in sekmeler():
        if ad not in wb.sheetnames:
            sekme_kur(wb, ad)
    return wb


def url_satiri(wb, url):
    for ws in wb.worksheets:
        for r in range(2, ws.max_row + 1):
            if ortak.ayni_adres(str(ws.cell(row=r, column=2).value or "").strip(), url):
                return ws, r
    return None, None


def kur(yol, hedef_dosya=None):
    import envanter
    wb = ac(yol)
    mevcut = wb.sheetnames
    var = {str(ws.cell(row=r, column=2).value).strip() for ws in wb.worksheets for r in range(2, ws.max_row + 1)
           if ws.cell(row=r, column=2).value}
    A = ortak.AYAR
    if hedef_dosya:
        urller = [x.strip() for x in open(hedef_dosya, encoding="utf-8") if x.strip() and not x.startswith("#")]
        sayfalar = []
        for u in urller:
            p = url_coz(u)
            if p:
                sayfalar.append(p)
            else:
                print(f"UYARI: adres çözülemedi, atlandı: {u}", file=sys.stderr)
    else:
        tipler = ("kategori", "cinsiyet_kategori") + (("marka", "marka_kategori", "marka_cinsiyet_kategori") if A.get("kur_marka") else ())
        haric = re.compile(A["kur_haric"]) if A.get("kur_haric") else None
        sayfalar = [s for s in envanter.yukle() if s["tip"] in tipler
                    and (not A.get("kur_kaynaklar") or s.get("kaynak") in A["kur_kaynaklar"])
                    and not (haric and (haric.search(s["url"]) or haric.search(s["slug"] or "")))]
    eklenen = {ad: 0 for ad in mevcut}
    for s in sorted(sayfalar, key=lambda s: s["slug"] or ""):
        if s["url"] in var:
            continue
        ws = wb[sekme_adi(s["url"], mevcut)]
        r = ws.max_row + 1
        # Ad slug'dan üretilir (Türkçe karaktersiz); brief hazırlanınca sayfanın gerçek adıyla değişir.
        hucre(ws, r, 1, (s["slug"] or "").replace("-", " ").title(), "kategori")
        hucre(ws, r, 2, s["url"], "url").hyperlink = s["url"]
        hucre(ws, r, 3, "Bekliyor", "brief")
        eklenen[ws.title] += 1
        var.add(s["url"])
    wb.save(yol)
    print(f"{yol} kuruldu · eklenen: " + ", ".join(f"{k} {v}" for k, v in eklenen.items()))


def ozet(yol):
    wb = ac(yol)
    for ws in wb.worksheets:
        d = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=2).value]
        print(f"{ws.title:10s} toplam {len(d):5d} · hazır {d.count('Hazır'):4d} · bekleyen {d.count('Bekliyor'):5d}")
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=3).value == "Hazır":
                print(f"   hazır: {ws.cell(row=r, column=1).value} · {ws.cell(row=r, column=2).value}")


def varsayilan_yol():
    A = ortak.AYAR
    ad = A.get("brief_dosyasi") or f"{A.get('marka_adi') or A.get('domain')} kategori içerik briefleri.xlsx"
    klasor = os.path.expanduser(A.get("cikti_klasoru") or ".")
    os.makedirs(klasor, exist_ok=True)
    return os.path.join(klasor, ad)


def main():
    from openpyxl.styles import PatternFill
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ortak.arguman_ekle(ap)
    ap.add_argument("--xlsx")
    ap.add_argument("--json")
    ap.add_argument("--kur", action="store_true")
    ap.add_argument("--hedefler", help="--kur ile: yalnız bu dosyadaki adresler eklenir (satır başına bir URL)")
    ap.add_argument("--ozet", action="store_true")
    a = ap.parse_args()
    d = json.load(open(a.json, encoding="utf-8")) if a.json else None
    if not a.marka and not a.domain and d:
        a.domain = d["url"]
    ortak.ayar_args(a)
    yol = a.xlsx or varsayilan_yol()
    if a.kur:
        return kur(yol, a.hedefler)
    if a.ozet:
        return ozet(yol)
    if not d:
        sys.exit("--json, --kur ya da --ozet gerekli")
    eksik = [k for k in ZORUNLU if k not in d]
    if eksik:
        sys.exit(f"satir.json içinde eksik alan: {', '.join(eksik)}")
    d["brief"] = "Hazır"

    wb = ac(yol)
    ws, r = url_satiri(wb, d["url"])
    guncelle = ws is not None
    if ws is None:
        ws = wb[sekme_adi(d["url"], wb.sheetnames)]
        r = ws.max_row + 1
    for c, k in enumerate(ANAHTAR, start=1):
        deger = d[k]
        if k == "kurgu" and d.get("kurgu_kalin"):
            kalin = str(d["kurgu_kalin"]).lstrip()
            if not str(d["kurgu"])[-1:].isspace():
                kalin = " " + kalin
            deger = CellRichText([TextBlock(InlineFont(rFont=FN, sz=10, b=False, color=INK), str(d["kurgu"])),
                                  TextBlock(InlineFont(rFont=FN, sz=10, b=True, color=INK), kalin)])
        cell = hucre(ws, r, c, deger, k)
        if k == "hacim":
            cell.number_format = "#,##0"
        if k == "url":
            cell.hyperlink = d["url"]
        if k == "brief":
            cell.fill = PatternFill("solid", fgColor="C8E6C9")
    en = max(satir_sayisi(str(d[k]) + str(d.get("kurgu_kalin", "") if k == "kurgu" else ""), GENISLIK[i])
             for i, k in enumerate(ANAHTAR))
    ws.row_dimensions[r].height = min(409, max(30, round(en * 13.2)))
    wb.save(yol)
    print(f"{yol} · sekme '{ws.title}' · {r}. satır · '{d['kategori']}' "
          f"{'dolduruldu' if guncelle else 'eklendi (envanterde yoktu)'}")
    if en * 13.2 > 409:
        print("UYARI: en uzun hücre 409 punto satır sınırını aşıyor (yaklaşık 30 satır); metnin tamamı ekranda "
              "görünmeyebilir. Kurgu satırları telgraf üslubuyla kısaltılabilir.")


if __name__ == "__main__":
    main()
