#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kategori içeriğini Word dosyasına basar (okuma ve onay kopyası).

Kullanım:
    python3 icerik_docx.py --json icerik.json --klasor "Kategori İçerik" [--marka X]

Belge adı "{slug}: {tam URL}" biçimindedir (ör. "kadin-mont: https://www.boyner.com.tr/kadin-mont-x-g3731-c23896554");
başlık satırına, belge özelliklerine ve dosya adına yazılır (dosya adında ":" ve "/" yerine görünüşü aynı olan
∶ ve ∕ karakterleri kullanılır).

Biçim marka profilinden gelir (markalar/{slug}/ayar.json; marka verilmezse adresin alan adından bulunur):
  liste: true   -> "mad" ve "li" gerçek Word listesi (madde imli / numaralı) olarak basılır
  liste: false  -> "mad" "•  " önekli, "li" "1. " önekli düz paragraf (CMS liste kabul etmiyorsa; ör. Boyner)
  tablo: true   -> ["tablo", [[başlık...], [satır...]]] Word tablosu olarak basılır
  tablo: false  -> içerikte tablo varsa belge üretilmez
  h1_govdede    -> gövdede H1 varsa belge başlığı olarak basılır, H2/H3 bir kademe aşağı iner

icerik.json biçimi:
{
  "kategori": "Kadın Mont",
  "url": "https://www.boyner.com.tr/kadin-mont-x-g3731-c23896554",
  "main_kw": "kadın mont",
  "meta": "İç künye; belgeye BASILMAZ",
  "linkler": {"LINK1": ["anchor metni", "https://www.boyner.com.tr/..."]},
  "govde": [["p","Başlıksız giriş paragrafı, içinde [LINK1] geçebilir"],
            ["H2","Başlık"], ["H3","Alt başlık"],
            ["mad","**Etiket:** madde satırı"],
            ["li","Sıralı adım"],
            ["tablo", [["Sütun", "Sütun"], ["hücre", "hücre [LINK2]"]]],      (yalnız profil tablo kabul ediyorsa)
            ["p","Vurgu için **kalın metin**"]],
  "sss": [["Soru?","Yanıt"]],
  "uzunluk": [750, 1250],         (isteğe bağlı: bu içerik için profil bandını geçersiz kılar; null = sınır yok)
  "sss_sayisi": [3, 4]            (isteğe bağlı: bu içerik için SSS soru sayısı bandı; null = araştırmada ne çıkarsa)
}
Son iki alan belgeye basılmaz; icerik_denetim.py okur.
Kaynak notu ve künye belgeye basılmaz; kullanıcıya sohbette söylenir.
"""
import argparse, json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ortak

INK, LINK, FN = "1A1A1A", "0B5FB0", "Calibri"


def renk(run, hex_kod):
    from docx.shared import RGBColor
    run.font.color.rgb = RGBColor.from_string(hex_kod)


def kopru(p, metin, url, boyut, kalin=False):
    from docx.oxml.shared import OxmlElement, qn
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                            is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    for etiket, deger in (("w:color", LINK), ("w:u", "single")):
        e = OxmlElement(etiket); e.set(qn("w:val"), deger); rPr.append(e)
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), FN); rf.set(qn("w:hAnsi"), FN); rPr.append(rf)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(boyut * 2))); rPr.append(sz)
    if kalin:
        rPr.append(OxmlElement("w:b"))
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = metin; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); p._p.append(h)


def metni_bas(p, metin, linkler, boyut=10.5, kalin_hepsi=False):
    from docx.shared import Pt
    for parca in re.split(r"(\[LINK\d+\]|\*\*[^*]+\*\*)", metin):
        m = re.fullmatch(r"\[(LINK\d+)\]", parca)
        if m and m.group(1) in linkler:
            kopru(p, linkler[m.group(1)][0], linkler[m.group(1)][1], boyut, kalin=kalin_hepsi)
        elif parca.startswith("**") and parca.endswith("**") and len(parca) > 4:
            # kalın parçanın içinde link olabilir: **[LINK1]:** -> kalın köprü + kalın ":"
            for alt in re.split(r"(\[LINK\d+\])", parca[2:-2]):
                ma = re.fullmatch(r"\[(LINK\d+)\]", alt)
                if ma and ma.group(1) in linkler:
                    kopru(p, linkler[ma.group(1)][0], linkler[ma.group(1)][1], boyut, kalin=True)
                elif alt:
                    r = p.add_run(alt); r.bold = True; r.font.name = FN; r.font.size = Pt(boyut); renk(r, INK)
        elif parca:
            r = p.add_run(parca); r.font.name = FN; r.font.size = Pt(boyut); renk(r, INK); r.bold = kalin_hepsi or None


def belge_adi(url):
    return ortak.slug_adi(url)


def dosya_adi(url):
    """Dosya adı da belge adıyla aynı görünür: "kadin-mont: https://www.boyner.com.tr/...".
    Dosya sistemleri adında ":" ve "/" kabul etmediği için görünüşü aynı olan Unicode karakterleri
    kullanılır: ∶ (U+2236) ve ∕ (U+2215). Belgenin içindeki başlık satırı ve belge özellikleri gerçek
    karakterleri taşır."""
    return f"{belge_adi(url)}: {url}".replace(":", "∶").replace("/", "∕") + ".docx"


def main():
    from docx import Document
    from docx.shared import Pt, Cm

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ortak.arguman_ekle(ap)
    ap.add_argument("--json", required=True)
    ap.add_argument("--out", help="çıktı yolu; verilmezse --klasor içine '{slug}: {URL}.docx'")
    ap.add_argument("--klasor", help="--out verilmediğinde dosyanın yazılacağı klasör (varsayılan: profildeki cikti_klasoru, yoksa .)")
    a = ap.parse_args()
    d = json.load(open(a.json, encoding="utf-8"))
    if not a.marka and not a.domain:
        a.domain = d["url"]
    A = ortak.ayar_args(a)
    klasor = os.path.expanduser(a.klasor or A.get("cikti_klasoru") or ".")
    os.makedirs(klasor, exist_ok=True)
    a.out = a.out or os.path.join(klasor, dosya_adi(d["url"]))
    linkler = {k: tuple(v) for k, v in (d.get("linkler") or {}).items()}
    liste_bicimi, tablo_izni = bool(A.get("liste")), bool(A.get("tablo"))
    if any(t == "tablo" for t, _ in d["govde"]) and not tablo_izni:
        raise SystemExit("İçerikte tablo var; profil tablo kabul etmiyor (tablo: false). Bilgiyi '•' satırlarına çevirin.")
    h1_var = any(t == "H1" for t, _ in d["govde"])

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FN; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2)

    # Belge adı: "{slug}: {tam URL}" (kullanıcı kararı 02.10.2026); başlık satırında ve belge özelliklerinde tam haliyle
    ad = belge_adi(d["url"])
    doc.core_properties.title = f"{ad}: {d['url']}"
    p = doc.add_paragraph(); r = p.add_run(f"{ad}: "); r.bold = True; r.font.size = Pt(13); r.font.name = FN; renk(r, INK)
    kopru(p, d["url"], d["url"], 13)

    def baslik(metin, seviye):
        # sayfada H1 varken (varsayılan) H2 -> Heading 1, H3 -> Heading 2; gövdede H1 varsa her biri kendi seviyesinde
        stil = {1: "Heading 1", 2: "Heading 2" if h1_var else "Heading 1", 3: "Heading 3" if h1_var else "Heading 2"}[seviye]
        p = doc.add_paragraph(style=stil)
        p.paragraph_format.space_before = Pt({1: 16, 2: 14, 3: 11}[seviye]); p.paragraph_format.space_after = Pt(5)
        r = p.add_run(metin); r.bold = True; r.font.name = FN; r.font.size = Pt({1: 16, 2: 13, 3: 11.5}[seviye]); renk(r, INK)

    def tablo_bas(satirlar):
        t = doc.add_table(rows=0, cols=max(len(s) for s in satirlar))
        t.style = "Table Grid"
        for i, s in enumerate(satirlar):
            hucreler = t.add_row().cells
            for j, deger in enumerate(s):
                hp = hucreler[j].paragraphs[0]
                metni_bas(hp, str(deger), linkler, 10, kalin_hepsi=(i == 0))
        doc.add_paragraph()

    sayac = 0
    for tip, icerik in d["govde"]:
        if tip != "li":
            sayac = 0
        if tip in ("H1", "H2", "H3"):
            baslik(icerik, int(tip[1]))
        elif tip == "tablo":
            tablo_bas(icerik)
        elif tip in ("li", "mad"):
            if liste_bicimi:
                p = doc.add_paragraph(style="List Number" if tip == "li" else "List Bullet")
                p.paragraph_format.space_after = Pt(4)
                metni_bas(p, icerik, linkler)
            else:
                # Liste biçimi kullanılmaz (CMS'e eklenemiyor): madde "•  " önekli, adım "1. " önekli düz paragraf.
                sayac += 1 if tip == "li" else 0
                onek = f"{sayac}. " if tip == "li" else "•  "
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(5)
                metni_bas(p, onek + icerik, linkler)
        else:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7); metni_bas(p, icerik, linkler)

    if d.get("sss"):
        baslik(d.get("sss_baslik") or f"{d['kategori']} Hakkında Sık Sorulan Sorular", 2)
        for soru, yanit in d["sss"]:
            baslik(soru, 3)
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); metni_bas(p, yanit, linkler)

    doc.save(a.out)
    duz = lambda i: re.sub(r"\*\*", "", re.sub(r"\[(LINK\d+)\]", lambda m: linkler.get(m.group(1), ("x",))[0], i))
    kelime = sum(len(duz(i).split()) for t, i in d["govde"] if t in ("p", "li", "mad"))
    sss_k = sum(len(duz(c).split()) for _, c in d.get("sss", []))
    print(f"yazıldı: {a.out} · gövde {kelime} kelime + SSS {sss_k} kelime · "
          f"{sum(1 for t, _ in d['govde'] if t == 'H2')} H2 · {sum(1 for t, _ in d['govde'] if t == 'H3')} H3 · "
          f"{len(linkler)} link · {len(d.get('sss', []))} SSS · biçim: {'Word listesi' if liste_bicimi else '• düz satır'}"
          f"{' · tablo' if any(t == 'tablo' for t, _ in d['govde']) else ''}")


if __name__ == "__main__":
    main()
